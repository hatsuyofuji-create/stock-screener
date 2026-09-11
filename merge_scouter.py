# -*- coding: utf-8 -*-
"""
=====================================================================
 銘柄スカウター（マネックス証券）のスクリーニング CSV を銘柄一覧表に流し込む
=====================================================================
 CSV は「10年スクリーニング」の CSV ダウンロード（Shift-JIS）。必要な列:
   コード, 市場, 業種, [指標]ROIC(%), [指標]PBR(倍), [指標]予想PER(会社予想)(倍),
   [指標]予想配当利回り(%), [指標]実績ROA(%)
 1回のダウンロードは 200 行までなので、複数ファイルをまとめて渡せる。

 表に入れる列:
   I ROIC / J PBR / K ROA      … CSV の値
   D PBR3以下 / E PER15以下 / F 配利3％以上 … CSV の値から 〇/✖ を判定（赤字予想は「－」）
   G 業種 / Q 市場             … CSV の値（東プ→プライム 等に変換）
   L ROIC順 / M PBR順 / N ROA順 / O 合成 … --rerank 指定時、全行に RANK 式を入れる

 上書きの考え方:
   既に値があるセルは触らない（既定）。
   --original 元の表.xlsx を渡すと「元の表で空欄だったセル」だけは CSV の値で上書きする
   （fill_table.py が J-Quants から入れた値を、銘柄スカウターの値に置き換えたいとき用）。

 使い方:
   python merge_scouter.py 銘柄一覧_filled.xlsx screening_1.csv screening_2.csv ... \
       -o 銘柄一覧_final.xlsx --original 銘柄一覧_補完.xlsx --rerank
=====================================================================
"""

from __future__ import annotations

import argparse
import copy
import csv
import math

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

# ============== 【設定】 ==============
PBR_MAX = 3.0
PER_MAX = 15.0
YIELD_MIN = 3.0
FIRST_ROW = 5
LAST_ROW = 740          # RANK 式の範囲の下端（表の最終行）
# =====================================

MARU, BATSU, DASH = "〇", "✖", "－"
MARKET = {"東プ": "プライム", "東ス": "スタンダード", "東グ": "グロース",
          "名メ": "名証メイン", "名ネ": "名証ネクスト", "札": "札証", "福": "福証"}
SECTOR = {"建設": "建設業", "電気・ガス": "電気・ガス業", "卸売": "卸売業", "小売": "小売業",
          "不動産": "不動産業", "情報・通信": "情報・通信業", "サービス": "サービス業",
          "鉄鋼": "鉄鋼", "非鉄": "非鉄金属", "金属": "金属製品", "化学": "化学",
          "医薬": "医薬品", "食料": "食料品", "繊維": "繊維製品", "パルプ・紙": "パルプ・紙",
          "ゴム": "ゴム製品", "ガラス・土石": "ガラス・土石製品", "石油・石炭": "石油・石炭製品",
          "陸運": "陸運業", "海運": "海運業", "空運": "空運業", "倉庫": "倉庫・運輸関連業",
          "銀行": "銀行業", "証券": "証券、商品先物取引業", "保険": "保険業",
          "その他金融": "その他金融業", "鉱業": "鉱業", "水産・農林": "水産・農林業",
          "その他製品": "その他製品"}
NO_FILL = PatternFill(fill_type=None)


def fnum(v):
    try:
        f = float(str(v).replace(",", ""))
        return None if math.isnan(f) else f
    except (TypeError, ValueError):
        return None


def col(headers: list[str], *keys: str) -> int | None:
    """見出しにキーワードをすべて含む列番号を返す。"""
    for i, h in enumerate(headers):
        if all(k in h for k in keys):
            return i
    return None


def read_csvs(paths: list[str]) -> dict[str, dict]:
    """複数 CSV を読んで コード -> {roic, pbr, per, yield, roa, market, sector} にまとめる。"""
    out: dict[str, dict] = {}
    for p in paths:
        with open(p, encoding="cp932", newline="") as f:
            rows = list(csv.reader(f))
        h = rows[0]
        ci = {"code": col(h, "コード"), "market": col(h, "市場"), "sector": col(h, "業種"),
              "roic": col(h, "ROIC"), "pbr": col(h, "PBR"), "per": col(h, "PER"),
              "yield": col(h, "配当利回り"), "roa": col(h, "ROA")}
        if ci["code"] is None:
            raise SystemExit(f"{p}: 「コード」列が見つかりません。見出し: {h}")
        for r in rows[1:]:
            if not r or not r[ci["code"]].strip():
                continue
            code = r[ci["code"]].strip()
            d = {}
            for k in ("roic", "pbr", "per", "yield", "roa"):
                d[k] = fnum(r[ci[k]]) if ci[k] is not None else None
            for k in ("market", "sector"):
                d[k] = r[ci[k]].strip() if ci[k] is not None else ""
            out[code] = d
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description="銘柄スカウターの CSV を銘柄一覧表に流し込む")
    ap.add_argument("xlsx")
    ap.add_argument("csv", nargs="+")
    ap.add_argument("-o", "--output", required=True)
    ap.add_argument("--original", help="元の表。ここで空欄だったセルは CSV の値で上書きする")
    ap.add_argument("--rerank", action="store_true", help="L/M/N/O 列に順位の式を入れる")
    args = ap.parse_args()

    data = read_csvs(args.csv)
    wb = load_workbook(args.xlsx)
    ws = wb.active
    orig = load_workbook(args.original).active if args.original else None

    ref_row = next((r for r in range(FIRST_ROW, ws.max_row + 1)
                    if all(ws.cell(r, c).value is not None for c in (9, 10, 11))), FIRST_ROW)

    def writable(r: int, c: str) -> bool:
        cell = ws[f"{c}{r}"]
        if cell.value is None:
            return True
        return orig is not None and orig[f"{c}{r}"].value is None

    def put(r: int, c: str, val) -> bool:
        if val is None or not writable(r, c):
            return False
        cell = ws[f"{c}{r}"]
        if not cell.has_style:
            cell._style = copy.copy(ws[f"{c}{ref_row}"]._style)
        cell.value = val
        cell.fill = NO_FILL  # 要確認の黄色を消す（CSV で確認済み）
        return True

    n_hit = 0
    for r in range(FIRST_ROW, ws.max_row + 1):
        code = ws.cell(r, 1).value
        if code is None or str(code).strip() not in data:
            continue
        d = data[str(code).strip()]
        n_hit += 1
        put(r, "I", d["roic"])
        put(r, "J", d["pbr"])
        put(r, "K", d["roa"])
        if d["pbr"] is not None:
            put(r, "D", MARU if d["pbr"] <= PBR_MAX else BATSU)
        if d["per"] is not None:
            put(r, "E", MARU if d["per"] <= PER_MAX else BATSU)
        elif d["roic"] is not None:  # 指標は取れているのに予想PERが無い = 赤字予想など
            put(r, "E", DASH)
        if d["yield"] is not None:
            put(r, "F", MARU if d["yield"] >= YIELD_MIN else BATSU)
        put(r, "Q", MARKET.get(d["market"], d["market"] or None))
        put(r, "G", SECTOR.get(d["sector"], d["sector"] or None))

    if args.rerank:
        rerank(ws)

    wb.calculation.fullCalcOnLoad = True
    wb.save(args.output)
    print(f"完了。CSV {len(data)} 銘柄のうち表にある {n_hit} 銘柄を反映し、{args.output} に保存しました。")


def rerank(ws) -> None:
    """L/M/N に RANK 式、O に合成式を入れる（値のある行だけ）。"""
    lo, hi = FIRST_ROW, LAST_ROW
    for r in range(lo, hi + 1):
        if ws.cell(r, 1).value is None:
            continue
        for c_from, c_to, order in (("I", "L", 0), ("J", "M", 1), ("K", "N", 0)):
            ws[f"{c_to}{r}"].value = (
                f'=IF(ISNUMBER({c_from}{r}),RANK({c_from}{r},${c_from}${lo}:${c_from}${hi},{order}),"")'
            )
        ws[f"O{r}"].value = f'=IF(AND(ISNUMBER(L{r}),ISNUMBER(M{r}),ISNUMBER(N{r})),L{r}+M{r}+N{r},"")'


if __name__ == "__main__":
    main()
