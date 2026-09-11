# -*- coding: utf-8 -*-
"""
=====================================================================
 銘柄一覧表（Excel）の空欄を J-Quants API で埋めるスクリプト
=====================================================================
 対象の表（Sheet1、5行目からデータ）:
   A 証券コード  B 銘柄名  C 下値限定的(半年)  D PBR3以下  E PER15以下
   F 配利3％以上  G 業種  H 詳細業種  I ROIC  J PBR  K ROA
   L ROIC順  M PBR順  N ROA順  O 合成  P 受注開示  Q 市場  R J-Quantsコード

 埋める列（空欄のセルだけ。既に値があるセルは触らない）:
   G 業種 / Q 市場 / R J-Quantsコード   … /equities/master
   J PBR / K ROA                        … /fins/summary ＋ 株価
   D E F の 〇/✖                         … PBR・PER・配当利回りから判定
   C 下値限定的(半年)                     … 直近6か月安値までの下落余地が 10% 以内なら〇
   M PBR順 / N ROA順 / O 合成             … --rerank 指定時のみ全行で再計算

 埋めない列（データが取れない・人の判断が要る）:
   H 詳細業種、I ROIC（有利子負債が summary に無い）、L ROIC順、P 受注開示

 使い方:
   1. このファイルと同じフォルダに .env というファイルを作り、1行書く:
        JQUANTS_API_KEY=ここにAPIキー
      （.env.example をコピーして書き換えればOK。.env は git に入らない）
   2. 実行:
        pip install openpyxl requests python-dotenv
        python fill_table.py 銘柄一覧.xlsx -o 銘柄一覧_filled.xlsx
        python fill_table.py 銘柄一覧.xlsx --provider mock   # 鍵なしで動作確認（乱数）
   ※ .env を使わず、環境変数 JQUANTS_API_KEY を設定しても動く。
   ※ レート制限(429)で途中で止まっても、出力ファイルを入力にして同じコマンドを
      もう一度実行すれば、埋まっていない行だけ取り直す（続きから再開できる）。
        python fill_table.py 銘柄一覧_filled.xlsx -o 銘柄一覧_filled.xlsx

 ※ V2 の項目名（短縮名）は環境により揺れがあるため、候補を FIELDS に並べてある。
   「フィールドが見つからない」というエラーが出たら、表示された実フィールド名を
   FIELDS に追加すればよい。
=====================================================================
"""

from __future__ import annotations

import argparse
import copy
import datetime as dt
import os
import random
import sys
import time

import requests
from openpyxl import load_workbook

# .env（このファイルと同じフォルダ）があれば読み込む。python-dotenv が無くても動く。
try:
    from dotenv import load_dotenv

    load_dotenv(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env"))
except ImportError:
    pass

# ============== 【設定】ここの数字だけ変えればOK ==============
PBR_MAX = 3.0        # D列: PBR がこれ以下なら〇
PER_MAX = 15.0       # E列: PER がこれ以下なら〇
YIELD_MIN = 3.0      # F列: 配当利回り(%) がこれ以上なら〇
DOWNSIDE_MAX = 10.0  # C列: 直近半年安値までの下落余地(%) がこれ以内なら〇
LOOKBACK_DAYS = 183  # C列で見る期間（暦日。約半年）
FIRST_ROW = 5        # データの開始行
SLEEP = 0.5          # API 呼び出し間隔（秒）
RETRY_WAIT = [5, 15, 30, 60, 90, 120, 180, 300]  # 429(レート制限) 時の待ち秒数（順に長く）
# ============================================================

MARU, BATSU, DASH = "〇", "✖", "－"

# V2（短縮名）と V1（旧名）の候補。最初に見つかった非空の値を使う。
FIELDS = {
    # --- /equities/master ---
    "code":     ["Code", "code", "LocalCode"],
    "sector33": ["S33Nm", "Sector33CodeName", "Sector33Name"],
    "market":   ["MktNm", "MarketCodeName", "MarketName", "MktCdNm"],
    # --- /equities/bars/daily（V2 実測: Date, C, L, AdjC, AdjL ...）---
    "date":     ["Date", "DiscDate", "DisclosedDate"],
    "close":    ["AdjC", "C", "Close", "AdjustmentClose"],
    "low":      ["AdjL", "L", "Low", "AdjustmentLow"],
    # --- /fins/summary（V2 実測: DiscDate, CurPerType, EPS, FEPS, NxFEPS, BPS, TA, NP,
    #     DivAnn, FDivAnn, NxFDivAnn ...）---
    "bps":      ["BPS", "BookValuePerShare"],
    "eps_fc":   ["FEPS", "NxFEPS", "ForecastEarningsPerShare",
                 "NextYearForecastEarningsPerShare"],   # 通期の予想EPS（FY開示なら来期予想）
    "eps":      ["EPS", "EarningsPerShare"],           # 実績EPS（通期開示のものを使う）
    "dps_fc":   ["FDivAnn", "NxFDivAnn", "ForecastDividendPerShareAnnual",
                 "NextYearForecastDividendPerShareAnnual"],  # 年間予想配当
    "dps":      ["DivAnn", "ResultDividendPerShareAnnual"],  # 年間実績配当
    "np":       ["NP", "Profit", "NetProfit"],
    "ta":       ["TA", "TotalAssets"],
    "period":   ["CurPerType", "TypeOfCurrentPeriod", "PeriodType"],
    "doc":      ["DocType", "TypeOfDocument"],
}


def first(row: dict, key: str):
    """FIELDS[key] の候補から最初に見つかった非空の値を返す。"""
    for n in FIELDS[key]:
        v = row.get(n)
        if v is not None and v != "":
            return v
    return None


def fnum(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def jq_code(code) -> str:
    """証券コード(4桁 or 3桁+英字) → J-Quants の5桁コード（末尾に0）。"""
    s = str(code).strip()
    return s + "0" if len(s) == 4 else s


# =========================================================
# データ取得（J-Quants V2 / mock）
# =========================================================
class JQuants:
    def __init__(self) -> None:
        key = os.getenv("JQUANTS_API_KEY", "").strip()
        if not key:
            raise SystemExit("JQUANTS_API_KEY が未設定です（J-Quants V2 の APIキー）。")
        self.base = (os.getenv("JQUANTS_BASE") or "https://api.jquants.com/v2").rstrip("/")
        self.s = requests.Session()
        self.s.headers["x-api-key"] = key

    def _get_once(self, path: str, params: dict) -> requests.Response:
        """1回 GET する。429(レート制限) は待って再試行する。"""
        for i, wait in enumerate(RETRY_WAIT + [None]):
            r = self.s.get(f"{self.base}{path}", params=params, timeout=30)
            if r.status_code != 429 or wait is None:
                return r
            ra = r.headers.get("Retry-After")
            if ra and ra.isdigit():
                wait = max(wait, int(ra))
            print(f"  レート制限にかかったので {wait} 秒待ちます（{i + 1}回目）", flush=True)
            time.sleep(wait)
        return r

    def get(self, path: str, params: dict) -> list[dict]:
        out, params = [], dict(params)
        while True:
            r = self._get_once(path, params)
            if r.status_code >= 400:
                raise RuntimeError(f"J-Quants {r.status_code} {path} {params}: {r.text[:300]}")
            body = r.json()
            out.extend(body.get("data", []))
            nxt = body.get("pagination_key")
            if not nxt:
                return out
            params["pagination_key"] = nxt

    def master(self) -> dict[str, dict]:
        rows = self.get("/equities/master", {})
        if not rows:
            raise RuntimeError("/equities/master が空でした。")
        m = {}
        for r in rows:
            c = first(r, "code")
            if c:
                m[str(c)] = {"sector": first(r, "sector33"), "market": first(r, "market")}
        if all(v["sector"] is None for v in m.values()):
            raise RuntimeError(f"業種フィールドが見つかりません。実フィールド例: {list(rows[0].keys())}")
        return m

    def bars(self, code: str, start: dt.date, end: dt.date) -> list[dict]:
        rows = self.get("/equities/bars/daily",
                        {"code": code, "from": start.isoformat(), "to": end.isoformat()})
        time.sleep(SLEEP)
        return rows

    def summary(self, code: str) -> list[dict]:
        rows = self.get("/fins/summary", {"code": code})
        time.sleep(SLEEP)
        return rows


class Mock:
    """鍵なしで動作確認するための乱数プロバイダ。"""

    def master(self):
        return {}

    def bars(self, code, start, end):
        random.seed(code)
        p = random.uniform(500, 5000)
        d, out = start, []
        while d <= end:
            if d.weekday() < 5:
                p *= random.uniform(0.98, 1.02)
                out.append({"Code": code, "Date": d.isoformat(), "C": round(p, 1),
                            "L": round(p * 0.99, 1)})
            d += dt.timedelta(days=1)
        return out

    def summary(self, code):
        random.seed(code + "f")
        bps = random.uniform(300, 4000)
        eps = random.uniform(-50, 400)
        return [{"Code": code, "DiscDate": "2026-05-10", "CurPerType": "FY",
                 "BPS": bps, "EPS": eps, "NxFEPS": eps * 1.05,
                 "NxFDivAnn": max(0.0, eps * 0.3), "NP": eps * 1e6, "TA": bps * 2.5e6},
                {"Code": code, "DiscDate": "2026-08-05", "CurPerType": "1Q",
                 "BPS": bps * 1.01, "EPS": eps / 4, "FEPS": eps * 1.05,
                 "FDivAnn": max(0.0, eps * 0.3), "NP": eps * 2.5e5, "TA": bps * 2.5e6}]


# =========================================================
# 指標の計算
# =========================================================
def latest_fy(rows: list[dict]) -> tuple[dict | None, dict | None]:
    """(直近の開示, 直近の通期決算) を返す。"""
    if not rows:
        return None, None
    rows = sorted(rows, key=lambda r: str(first(r, "date") or ""))
    latest = rows[-1]
    fy = [r for r in rows if str(first(r, "period") or "").upper() == "FY"]
    return latest, (fy[-1] if fy else None)


def compute(code: str, prov, today: dt.date) -> dict:
    """1銘柄ぶんの指標を dict で返す。取れない項目は None。"""
    out = {"price": None, "low6m": None, "pbr": None, "per": None, "yield": None,
           "roa": None}
    bars = prov.bars(code, today - dt.timedelta(days=LOOKBACK_DAYS), today)
    bars = sorted(bars, key=lambda b: str(first(b, "date") or ""))  # 古い→新しい
    closes = [fnum(first(b, "close")) for b in bars]
    lows = [fnum(first(b, "low")) for b in bars]
    closes = [c for c in closes if c]
    lows = [l for l in lows if l]
    if closes:
        out["price"] = closes[-1]
        out["low6m"] = min(lows or closes)

    latest, fy = latest_fy(prov.summary(code))
    price = out["price"]
    if latest and price:
        # 予想EPS・予想配当は直近開示のものを優先し、無ければ通期開示の実績を使う
        # （四半期開示の EPS は四半期ぶんの値なので PER には使わない）
        bps = fnum(first(latest, "bps")) or (fy and fnum(first(fy, "bps")))
        eps = fnum(first(latest, "eps_fc"))
        if eps is None and fy:
            eps = fnum(first(fy, "eps_fc")) or fnum(first(fy, "eps"))
        dps = fnum(first(latest, "dps_fc"))
        if dps is None and fy:
            dps = fnum(first(fy, "dps_fc")) or fnum(first(fy, "dps"))
        if bps and bps > 0:
            out["pbr"] = price / bps
        if eps and eps > 0:
            out["per"] = price / eps
        elif eps is not None:
            out["per"] = -1.0  # 赤字予想 → 「－」
        if dps is not None:
            out["yield"] = dps / price * 100.0
    src = fy or latest
    if src:
        np_, ta = fnum(first(src, "np")), fnum(first(src, "ta"))
        if np_ is not None and ta:
            out["roa"] = np_ / ta * 100.0
    return out


def flag(value, ok) -> str:
    if value is None:
        return DASH
    return MARU if ok(value) else BATSU


# =========================================================
# Excel の読み書き
# =========================================================
def main() -> None:
    ap = argparse.ArgumentParser(description="銘柄一覧表の空欄を J-Quants で埋める")
    ap.add_argument("xlsx", nargs="?", help="入力の Excel ファイル（--show-fields のときは不要）")
    ap.add_argument("-o", "--output", help="出力ファイル（省略時は *_filled.xlsx）")
    ap.add_argument("--provider", choices=["jquants", "mock"], default="jquants")
    ap.add_argument("--rerank", action="store_true",
                    help="M列(PBR順)・N列(ROA順)・O列(合成)を、値のある全行で再計算する")
    ap.add_argument("--limit", type=int, default=0, help="先頭 N 銘柄だけ処理（動作確認用）")
    ap.add_argument("--show-fields", metavar="CODE",
                    help="指定コード(例: 6307)の API 応答の項目名と値を表示して終了（項目名の確認用）")
    args = ap.parse_args()

    if args.show_fields:
        show_fields(args.show_fields)
        return
    if not args.xlsx:
        ap.error("入力の Excel ファイルを指定してください")

    out_path = args.output or args.xlsx.rsplit(".", 1)[0] + "_filled.xlsx"
    prov = Mock() if args.provider == "mock" else JQuants()
    today = dt.date.today()

    wb = load_workbook(args.xlsx)
    ws = wb.active
    master = prov.master()

    # 書式の手本: R列まで埋まっている最初の行
    ref_row = next((r for r in range(FIRST_ROW, ws.max_row + 1)
                    if all(ws.cell(r, c).value is not None for c in (1, 7, 9, 17, 18))),
                   FIRST_ROW)

    def put(r: int, col: str, val) -> None:
        cell = ws[f"{col}{r}"]
        if cell.value is not None:
            return  # 既存の値は上書きしない
        if not cell.has_style:
            cell._style = copy.copy(ws[f"{col}{ref_row}"]._style)
        cell.value = val

    n_done = n_skip = n_fail = 0
    for r in range(FIRST_ROW, ws.max_row + 1):
        code = ws.cell(r, 1).value
        if code is None:
            continue
        code5 = jq_code(code)
        put(r, "R", int(code5) if code5.isdigit() else code5)

        # 埋める対象の列がすべて埋まっている行は通信しない（再実行時の続きから用）
        if all(ws[f"{c}{r}"].value is not None for c in "CDEFJK"):
            n_skip += 1
            continue

        m = master.get(code5)
        if m:
            if m["sector"]:
                put(r, "G", m["sector"])
            if m["market"]:
                put(r, "Q", str(m["market"]).replace("市場", ""))

        try:
            v = compute(code5, prov, today)
        except Exception as e:  # 1銘柄の失敗で全体を止めない
            n_fail += 1
            print(f"  ! {code}: {e}", file=sys.stderr)
            if n_fail >= 20:
                print("  失敗が続くので中断します。ここまでの結果は保存します。"
                      " しばらく待ってから同じコマンドをもう一度実行すると続きから埋まります。")
                break
            continue

        if v["pbr"] is not None:
            put(r, "J", round(v["pbr"], 2))
            put(r, "D", flag(v["pbr"], lambda x: x <= PBR_MAX))
        if v["per"] is not None:
            put(r, "E", DASH if v["per"] < 0 else flag(v["per"], lambda x: x <= PER_MAX))
        if v["yield"] is not None:
            put(r, "F", flag(v["yield"], lambda x: x >= YIELD_MIN))
        if v["roa"] is not None:
            put(r, "K", round(v["roa"], 2))
        if v["price"] and v["low6m"]:
            downside = (v["price"] - v["low6m"]) / v["price"] * 100.0
            put(r, "C", flag(downside, lambda x: x <= DOWNSIDE_MAX))

        n_done += 1
        if n_done % 50 == 0:
            print(f"  ...{n_done} 銘柄 処理済み")
        if args.limit and n_done >= args.limit:
            break

    if args.rerank:
        rerank(ws)

    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)
    print(f"完了。取得 {n_done} 銘柄 / 埋まっていたので省略 {n_skip} 銘柄 / 失敗 {n_fail} 銘柄"
          f" → {out_path} に保存しました。")
    if n_fail:
        print("失敗した銘柄は空欄のままです。同じコマンドをもう一度実行すると、そこだけ取り直します。")


def show_fields(code) -> None:
    """/fins/summary と /equities/bars/daily の実際の項目名を表示する（FIELDS 調整用）。"""
    prov = JQuants()
    code5 = jq_code(code)
    today = dt.date.today()
    print(f"=== /equities/bars/daily {code5}（直近1件） ===")
    bars = prov.bars(code5, today - dt.timedelta(days=14), today)
    for k, v in (bars[-1] if bars else {}).items():
        print(f"  {k} = {v}")
    print(f"=== /fins/summary {code5}（直近1件） ===")
    latest, _ = latest_fy(prov.summary(code5))
    for k, v in (latest or {}).items():
        print(f"  {k} = {v}")


def rerank(ws) -> None:
    """J列(PBR 小→大)・K列(ROA 大→小)の順位を値のある全行で振り直し、O列に合成式を入れる。"""
    rows = [r for r in range(FIRST_ROW, ws.max_row + 1) if ws.cell(r, 1).value is not None]
    pbr = sorted((r for r in rows if isinstance(ws.cell(r, 10).value, (int, float))),
                 key=lambda r: ws.cell(r, 10).value)
    roa = sorted((r for r in rows if isinstance(ws.cell(r, 11).value, (int, float))),
                 key=lambda r: -ws.cell(r, 11).value)
    for i, r in enumerate(pbr, 1):
        ws.cell(r, 13).value = i
    for i, r in enumerate(roa, 1):
        ws.cell(r, 14).value = i
    for r in rows:
        if all(isinstance(ws.cell(r, c).value, (int, float)) for c in (12, 13, 14)):
            ws.cell(r, 15).value = f"=L{r}+M{r}+N{r}"


if __name__ == "__main__":
    main()
